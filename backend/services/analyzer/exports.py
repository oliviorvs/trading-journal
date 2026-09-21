"""Exports — JSON, CSV, XLSX.

Le journal n'exporte aujourd'hui qu'en PDF : lisible, mais impossible à
retravailler. Ces trois formats répondent à « puis-je exporter mes stats vers
Excel ou les partager ? ».

Aucune nouvelle dépendance : `openpyxl` est DÉJÀ dans requirements.txt (il
sert à lire les rapports MT5 en xlsx lors de l'import). `csv` et `json` sont
dans la bibliothèque standard.
"""
from __future__ import annotations

import csv
import io
import json
from typing import Dict, List, Sequence

# Colonnes exportées pour une ligne de dimension, dans l'ordre. Le libellé est
# celui affiché à l'écran : un CSV dont les en-têtes diffèrent de l'interface
# oblige à deviner la correspondance.
DIMENSION_COLUMNS = [
    ("key", "Valeur"),
    ("trades", "Trades"),
    ("pct_of_total", "Part (%)"),
    ("win_rate", "Win rate (%)"),
    ("wins", "Gagnants"),
    ("losses", "Perdants"),
    ("breakeven", "Breakeven"),
    ("pnl", "Résultat"),
    ("avg_pnl", "Résultat moyen"),
    ("expectancy", "Expectancy"),
    ("expectancy_r", "Expectancy (R)"),
    ("avg_r", "R moyen"),
    ("r_coverage", "Couverture R (%)"),
    ("profit_factor", "Profit factor"),
    ("drawdown_money", "Drawdown cumulé"),
    ("status", "Fiabilité"),
]


def to_json(payload: dict) -> bytes:
    return json.dumps(payload, ensure_ascii=False, indent=2, default=str).encode("utf-8")


def _table_from_rows(rows: Sequence[dict], key_label: str = "Valeur") -> List[List]:
    header = [label if field != "key" else key_label for field, label in DIMENSION_COLUMNS]
    table = [header]
    for row in rows:
        table.append([_cell(row.get(field)) for field, _label in DIMENSION_COLUMNS])
    return table


def _cell(value):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    return value


def build_tables(bundle: dict) -> Dict[str, List[List]]:
    """Une table par dimension, à partir du paquet d'analyses complet."""
    tables: Dict[str, List[List]] = {}

    overview = bundle.get("overview") or {}
    analyzer = overview.get("analyzer") or {}
    journal = overview.get("journal") or {}
    tables["Vue d'ensemble"] = [["Indicateur", "Valeur"]] + [
        ["Trades", analyzer.get("trades")],
        ["Win rate (%)", analyzer.get("win_rate")],
        ["Résultat", analyzer.get("pnl")],
        ["Profit factor", _cell(analyzer.get("profit_factor"))],
        ["Expectancy", analyzer.get("expectancy")],
        ["Expectancy (R)", _cell(analyzer.get("expectancy_r"))],
        ["Couverture R (%)", analyzer.get("r_coverage")],
        ["Drawdown max (%)", journal.get("max_drawdown")],
        ["Drawdown max", journal.get("max_drawdown_abs")],
        ["Résultat brut", analyzer.get("gross_pnl")],
        ["Résultat net", analyzer.get("net_pnl")],
    ]

    for name, block in (bundle.get("dimensions") or {}).items():
        tables[_sheet_name(name)] = _table_from_rows(block.get("rows", []), _sheet_name(name))

    errors = bundle.get("errors") or {}
    if errors.get("rows"):
        table = _table_from_rows(errors["rows"], "Erreur")
        table[0] += ["Coût", "Coût (R)", "Δ expectancy"]
        for index, row in enumerate(errors["rows"], start=1):
            table[index] += [_cell(row.get("cost")), _cell(row.get("cost_r")),
                             _cell(row.get("delta_expectancy"))]
        tables["Coût des erreurs"] = table

    patterns = bundle.get("patterns") or {}
    if patterns.get("rows"):
        header = ["Condition", "n", "Win rate (%)", "IC bas", "IC haut", "Expectancy",
                  "Expectancy (R)", "Profit factor", "Drawdown cumulé", "Fiabilité", "Non concluant"]
        table = [header]
        for row in patterns["rows"]:
            ci = row.get("win_rate_ci") or {}
            table.append([
                row.get("condition"), row.get("trades"), row.get("win_rate"),
                _cell(ci.get("low")), _cell(ci.get("high")),
                row.get("expectancy"), _cell(row.get("expectancy_r")),
                _cell(row.get("profit_factor")), row.get("drawdown_money"),
                row.get("status"), _cell(row.get("inconclusive")),
            ])
        tables["Patterns"] = table

    quality = bundle.get("data_quality") or {}
    if quality.get("coverage"):
        table = [["Champ", "Renseignés", "Manquants", "Couverture (%)"]]
        for row in quality["coverage"]:
            table.append([row["field"], row["filled"], row["missing"], row["coverage"]])
        tables["Qualité des données"] = table

    insights = bundle.get("insights") or []
    if insights:
        tables["Constats automatiques"] = [["Type", "Titre", "Constat", "n"]] + [
            [item["kind"], item["title"], item["text"], item.get("trades")]
            for item in insights
        ]

    return tables


def _sheet_name(name: str) -> str:
    labels = {
        "setup": "Setups", "symbol": "Instruments", "session": "Sessions",
        "weekday": "Jours", "hour": "Heures", "month": "Mois", "week": "Semaines",
        "emotion": "Émotions", "error": "Erreurs", "direction": "Sens",
        "timeframe": "Timeframes", "exit_reason": "Sorties", "playbook": "Playbooks",
        "duration": "Durées", "tag": "Étiquettes", "risk_source": "Source du risque",
    }
    return labels.get(name, name.capitalize())


def to_csv(bundle: dict) -> bytes:
    """CSV unique, sections séparées — UTF-8 avec BOM et séparateur `;`.

    Ce choix n'est pas cosmétique : sans BOM, Excel sous Windows lit les
    accents en mojibake ; avec une virgule, il colle tout dans une seule
    colonne dans les locales francophones.
    """
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", lineterminator="\r\n")
    for title, table in build_tables(bundle).items():
        writer.writerow([f"### {title}"])
        writer.writerows(table)
        writer.writerow([])
    return b"\xef\xbb\xbf" + buffer.getvalue().encode("utf-8")


def to_xlsx(bundle: dict) -> bytes:
    """Une feuille par dimension (openpyxl, déjà présent)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="2F3B52")

    for title, table in build_tables(bundle).items():
        # Excel : 31 caractères maximum, et []:*?/\ interdits dans un nom.
        safe = title[:31]
        for char in "[]:*?/\\":
            safe = safe.replace(char, "-")
        sheet = workbook.create_sheet(safe)
        for row in table:
            sheet.append(row)
        for cell in sheet[1]:
            cell.font = header_font
            cell.fill = header_fill
        for column_cells in sheet.columns:
            longest = max((len(str(c.value)) for c in column_cells if c.value is not None), default=8)
            sheet.column_dimensions[column_cells[0].column_letter].width = min(max(longest + 2, 10), 40)
        sheet.freeze_panes = "A2"

    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()
